"""One-time, source-backed insert into the current user-edited DOCX.

Never rebuild the original report. The baseline must be backed up first.
"""
from pathlib import Path
import hashlib
import json
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
from build_task007_feishu_docx import add_inline
from update_task016_feishu_docx import after, table, picture

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / 'output/task-007/面向光子芯片智能设计的三方向技术调研_飞书上传版.docx'
BASE = ROOT / 'tmp/task-017/report-user-baseline.docx'
SOURCE = ROOT / 'references/task-007/candidate-research-direction.md'
FIG = ROOT / 'references/task-017-candidate-direction/figures'
MARKER = '6.6 当前候选主线：物理驱动的自主光子芯片设计'


def diagram(number):
    """Deterministic conceptual diagrams; no experimental data."""
    FIG.mkdir(parents=True, exist_ok=True)
    rows = {
        1: ['器件物理知识 + 真实仿真反馈', '物理感知推理 → 设计诊断', '设计决策 → Harness 执行 → 闭环 PIC 设计'],
        2: ['需求 / DesignSpec → 规划 → 初始设计', 'Solver / Optimizer → 仿真结果与条件记录',
            '【研究集中区域】物理诊断 → 设计决策 → 下一步动作',
            '下一步动作 → 再仿真；目标检查 → 版图 / 分层验证'],
        3: ['概念比较：不是成熟度排名', '传统优化：变量 / 目标搜索；逆设计：目标驱动结构搜索',
            '脚本生成 Agent：输出代码；工具调用 Agent：执行与反馈',
            '物理引导自主 Agent：拟研究诊断与决策（增益 UNKNOWN）'],
    }[number]
    im = Image.new('RGB', (1600, 100 + len(rows) * 115), 'white')
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 32)
    for i, text in enumerate(rows):
        y = 35 + i * 115
        draw.rounded_rectangle((30, y, 1570, y + 86), radius=10,
                               fill='#d6eee7' if number == 2 and i == 2 else '#e9f0f8')
        draw.text((800, y+43), text, font=font, anchor='mm', fill='#19383b')
        if i < len(rows)-1 and number != 3:
            draw.line((800, y+87, 800, y+113), fill='#167a73', width=3)
    p = FIG / f'concept-{number}.png'
    im.save(p)
    return p


def fingerprint(doc):
    # Full original XML blocks retained, including formatting and hyperlinks.
    from lxml import etree
    return [etree.tostring(x, method='c14n') for x in doc._element.body if not x.tag.endswith('sectPr')]


def main():
    if not BASE.exists():
        raise SystemExit('Missing user baseline backup')
    doc = Document(TARGET)
    if any(MARKER in p.text for p in doc.paragraphs):
        raise SystemExit('Refusing duplicate insert')
    if TARGET.read_bytes() != BASE.read_bytes():
        raise SystemExit('Current DOCX changed after backup; stop and rebase on user edits')
    before = fingerprint(doc)
    anchors = [p for p in doc.paragraphs if p.text.strip() == '参考文献与证据入口']
    if len(anchors) != 1:
        raise SystemExit('Reference heading anchor is not unique')
    cursor = anchors[0]._p.getprevious()

    def para(text, style='Normal'):
        nonlocal cursor
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.16
        if style.startswith('Heading'):
            p.paragraph_format.keep_with_next = True
        add_inline(p, text, size=11 if style.startswith('Heading') else 9.6)
        cursor = after(cursor, p._p)

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    para(MARKER, 'Heading 2')
    i, fig = 1, 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith('```mermaid'):
            fig += 1
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                i += 1
            cursor = picture(doc, cursor, diagram(fig), 6.4,
                f'候选方向概念图 {fig}（INFERENCE / USER RESEARCH PREFERENCE；非已实现系统或实验结果）')
        elif s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x) for x in row):
                    rows.append(row)
                i += 1
            widths = [6.4 / len(rows[0])] * len(rows[0])
            cursor = table(doc, cursor, rows, widths)
            i -= 1
        elif s.startswith('## '):
            m = re.match(r'## (\d+)\. (.*)', s)
            para(f'6.6.{m[1]} {m[2]}' if m else s[3:], 'Heading 3')
        elif s.startswith('### '):
            para(s[4:], 'Heading 4')
        elif s:
            # Relative source links are represented by actual web citations in
            # this section and a provenance paragraph, not broken Word paths.
            s = re.sub(r'\[([^\]]+)\]\((?!https?://)[^)]+\)', r'\1（见仓库同名证据材料）', s)
            para(s)
        i += 1
    para('本节事实源：references/task-007/candidate-research-direction.md；2026-09-30 用户候选主线授权。新增内容仅为候选分析，原有研究与用户编辑保留。')
    doc.save(TARGET)

    result = Document(TARGET)
    after_blocks = fingerprint(result)
    j = 0
    for b in after_blocks:
        if j < len(before) and b == before[j]:
            j += 1
    if j != len(before):
        raise SystemExit(f'Original XML block preservation failed: {j}/{len(before)}')
    with ZipFile(BASE) as a, ZipFile(TARGET) as b:
        original_media = [n for n in a.namelist() if n.startswith('word/media/') and not n.endswith('/')]
        assert all(a.read(n) == b.read(n) for n in original_media)
        for n in ['word/styles.xml', 'word/settings.xml']:
            from lxml import etree
            assert etree.tostring(etree.fromstring(a.read(n)), method='c14n') == etree.tostring(etree.fromstring(b.read(n)), method='c14n'), n
    qa = {'original_body_blocks_preserved': j, 'original_media_preserved':len(original_media),
          'tables_before':len(Document(BASE).tables), 'tables_after':len(result.tables),
          'baseline_sha256':hashlib.sha256(BASE.read_bytes()).hexdigest(),
          'output_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),
          'user_baseline_backup':str(BASE)}
    out = ROOT / 'references/task-017-candidate-direction'
    (out / 'docx-preservation.json').write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(qa, ensure_ascii=False))


if __name__ == '__main__':
    main()
