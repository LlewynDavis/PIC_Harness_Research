"""Package the accepted TASK-007 Markdown report as a self-contained DOCX.

The research text remains in references/task-007/formal-research-report.md.
This script only changes presentation and adds the existing source index as an
appendix so a Feishu upload does not depend on local MkDocs routes.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "references/task-007/formal-research-report.md"
INDEX = ROOT / "references/task-006/paper-and-repository-index.md"
OUTPUT = ROOT / "output/task-007/面向光子芯片智能设计的三方向技术调研_飞书上传版.docx"
ASSET_DIR = ROOT / "tmp/task-007-docx-qa"
ACCEPTED = "3ed876f5fce34d38a47907d6c68c476f6a51d090"
BASE = f"https://github.com/LlewynDavis/PIC_Harness_Research/blob/{ACCEPTED}/"
PORTAL_PATHS = {
    "/task-007-search.html": "references/task-007/search-and-screening-log.md",
    "/task-007-evidence-audit.html": "references/task-007/full-text-evidence-audit.md",
    "/task-007-terms.html": "references/task-007/terminology-ledger.md",
    "/task-007-routes.html": "references/task-007/candidate-routes-and-open-questions.md",
    "/task-006-source-index.html": "references/task-006/paper-and-repository-index.md",
    "/task-006-evidence.html": "references/task-006/source-log.md",
    "/technical-map-and-objects.html": "references/task-005/technical-map-and-objects.md",
    "/capability-and-verification-matrix.html": "references/task-005/capability-and-verification-matrix.md",
    "/reference-architecture-v0.html": "references/task-005/reference-architecture-v0.md",
}
FONT = "Microsoft YaHei"


def font_run(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)


def set_cell_margin(cell, top=95, start=95, bottom=95, end=95):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    item = OxmlElement("w:hyperlink")
    item.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "205493")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label.replace("`", "")
    run.append(text)
    item.append(run)
    paragraph._p.append(item)


TOKEN = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(paragraph, value, size=None):
    for part in TOKEN.split(value):
        if not part:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if match:
            url = PORTAL_PATHS.get(match.group(2), match.group(2))
            if url.startswith("references/"):
                url = BASE + url
            add_hyperlink(paragraph, match.group(1), url)
        elif part.startswith("**") and part.endswith("**"):
            font_run(paragraph.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = font_run(paragraph.add_run(part[1:-1]), size=size)
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = font_run(paragraph.add_run(part[1:-1]), size=size, color="44546A")
        else:
            font_run(paragraph.add_run(part), size=size)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    for name, size, before, after, bold in (
        ("Normal", 10.5, 0, 7, False),
        ("Title", 20, 0, 10, True),
        ("Subtitle", 12, 0, 9, False),
        ("Heading 1", 15, 16, 8, True),
        ("Heading 2", 12, 11, 6, True),
        ("Heading 3", 10.5, 9, 4, True),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_after = Pt(6)
    doc.styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    doc.core_properties.title = "面向光子芯片智能设计的三方向技术调研"
    doc.core_properties.subject = "TASK-007 正式技术调研报告"
    doc.core_properties.author = "PIC_Harness_Research"


def add_table(doc, rows, landscape=False):
    if not rows:
        return
    n = len(rows[0])
    if landscape:
        sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
        sec.left_margin = sec.right_margin = Inches(0.5)
        sec.top_margin = sec.bottom_margin = Inches(0.58)
        doc.add_heading("4.3 Harness Capability Matrix", level=2)
    table = doc.add_table(rows=1, cols=n)
    table.autofit = False
    width = 9.95 if landscape else 6.86
    if n == 9:
        ratios = [1.2, 0.9, 0.75, 0.85, 1.1, 1.25, 1.1, 1.25, 1.4]
    elif n == 4:
        ratios = [0.82, 1.25, 1.56, 1.5]
    elif n == 3:
        ratios = [1.04, 1.02, 1.56]
    else:
        ratios = [1] * n
    unit = width / sum(ratios)
    for r_idx, values in enumerate(rows):
        row = table.rows[0] if r_idx == 0 else table.add_row()
        cells = row.cells
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        if r_idx == 0:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(hdr)
        for i, value in enumerate(values):
            cell = cells[i]
            cell.width = Inches(unit * ratios[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell, 75 if n == 8 else 95, 85, 75 if n == 8 else 95, 85)
            if r_idx == 0:
                set_cell_shading(cell, "274766")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F4F7FA")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.16
            if n >= 8 and i in (1, 2, 3):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(para, value.strip(), size=7.7 if n >= 8 else 8.7)
            if r_idx == 0:
                for run in para.runs:
                    font_run(run, bold=True, color="FFFFFF")
    if landscape:
        sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.72)
        sec.left_margin = sec.right_margin = Inches(0.82)


def draw_arrow(draw, a, b, color="#6B7E90", width=4):
    draw.line((a, b), fill=color, width=width)
    x1, y1 = a
    x2, y2 = b
    if abs(x2 - x1) > abs(y2 - y1):
        d = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 12 * d, y2 - 7), (x2 - 12 * d, y2 + 7)]
    else:
        d = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 7, y2 - 12 * d), (x2 + 7, y2 - 12 * d)]
    draw.polygon(points, fill=color)


def make_diagram(kind, path):
    if kind == "map":
        labels = ["需求 / 自然语言", "DesignSpec / typed IR / DSL", "Agent / Harness / Planning", "Tool Registry / Adapter", "器件或电路设计", "物理 / 电路仿真", "优化 / 逆向设计", "版图 / GDS / Routing", "DRC / LVS / Physical Checks", "RESULT / Evaluator", "状态 / 失败恢复 / 人工升级"]
        side = ("Knowledge / RAG / PDK", 2)
    else:
        labels = ["Requirement", "DesignSpec / Direct Typed Contract", "Planner / Harness", "Tool Registry / Adapter", "Execution / Solver / CAD / EDA", "RESULT + Raw Artifacts", "Independent Evaluator", "State + Evidence Ledger", "Decision", "Layout / Verification Gate"]
        side = ("Knowledge / PDK / Provenance", 2)
    w, h = 1300, 1640
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 29)
    small = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 24)
    x0, x1 = 155, 825
    top, step, boxh = 45, 130 if kind == "map" else 138, 74
    if kind == "architecture":
        top = 34
    for i, label in enumerate(labels):
        y = top + i * step
        d.rounded_rectangle((x0, y, x1, y + boxh), radius=13, fill="#EEF3F8", outline="#8BA0B4", width=3)
        bb = d.textbbox((0, 0), label, font=font)
        tw = bb[2] - bb[0]
        d.text(((x0 + x1 - tw) / 2, y + 18), label, font=font, fill="#1E3348")
        if i:
            draw_arrow(d, ((x0 + x1) // 2, y - step + boxh + 4), ((x0 + x1) // 2, y - 6))
    sx0, sx1 = 930, 1250
    sy = top + side[1] * step
    d.rounded_rectangle((sx0, sy, sx1, sy + boxh + 22), radius=13, fill="#EAF4F3", outline="#6B9A96", width=3)
    side_lines = side[0].split(" / ")
    for j, t in enumerate(side_lines):
        bb = d.textbbox((0, 0), t, font=small)
        tw = bb[2] - bb[0]
        d.text(((sx0 + sx1 - tw) / 2, sy + 10 + j * 26), t, font=small, fill="#1E4D4A")
    draw_arrow(d, (sx0 - 8, sy + boxh // 2), (x1 + 8, sy + boxh // 2))
    if kind == "map":
        d.text((920, top + 8 * step), "反馈回规划", font=small, fill="#546A7D")
        d.line([(x1 + 12, top + 10 * step + boxh // 2), (890, top + 10 * step + boxh // 2), (890, top + 2 * step + boxh // 2), (x1 + 10, top + 2 * step + boxh // 2)], fill="#8AA0B2", width=3)
    else:
        by = top + 9 * step + boxh + 34
        d.text((140, by), "可恢复失败 → 返回规划", font=small, fill="#50677A")
        d.text((690, by), "不确定或高风险 → 人工升级", font=small, fill="#50677A")
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, dpi=(180, 180))


def add_diagram(doc, kind):
    path = ASSET_DIR / f"{kind}.png"
    make_diagram(kind, path)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(7)
    para.add_run().add_picture(str(path), width=Inches(6.65))


def parse_table(lines):
    rows = []
    for line in lines:
        values = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", v) for v in values):
            continue
        rows.append(values)
    return rows


def append_source_index(doc):
    doc.add_heading("附录 C 论文与项目出处索引", level=1)
    p = doc.add_paragraph()
    add_inline(p, "以下条目依据已验收的论文与项目出处索引整理，提供可点击的一手论文和仓库入口。出处链接不意味着本项目复算或复现作者结果。")
    text = INDEX.read_text(encoding="utf-8").splitlines()
    wanted = {
        "面向智能体的 PIC 设计与方法论文",
        "DAS 光电混合集成论文",
        "光子人工智能计算论文",
        "纯项目与工具出处",
    }
    section = None
    items = 0
    for line in text:
        if line.startswith("## "):
            section = line[3:].strip() if line[3:].strip() in wanted else None
            if section:
                doc.add_heading(section, level=2)
            continue
        if not section or not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if not cells or cells[0] in ("对象", "项目") or all(re.fullmatch(r":?-{3,}:?", v) for v in cells):
            continue
        items += 1
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, f"{items}. " + " · ".join(cells), size=8.8)
    return items


def build():
    doc = Document()
    setup_styles(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    tables = 0
    diagrams = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = parse_table(block)
            add_table(doc, rows, landscape=len(rows[0]) >= 7)
            tables += 1
            continue
        if line.startswith("```"):
            typ = line[3:]
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            if typ == "mermaid":
                add_diagram(doc, "map" if diagrams == 0 else "architecture")
                diagrams += 1
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_after = Pt(9)
                for n, item in enumerate(block):
                    if n:
                        p.add_run().add_break()
                    font_run(p.add_run(item), size=9.6, color="34495E")
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            label = line[level:].strip()
            if i == 0:
                para = doc.add_paragraph(style="Title")
                add_inline(para, label)
            elif i == 2:
                para = doc.add_paragraph(style="Subtitle")
                add_inline(para, label)
            elif label == "4.3 Harness Capability Matrix":
                i += 1
                continue
            else:
                para = doc.add_heading(level=min(level, 3))
                add_inline(para, label)
            i += 1
            continue
        if line.startswith("**项目**") or line.startswith("**任务**") or line.startswith("**证据截止**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line, size=9.5)
            i += 1
            continue
        if line.startswith("**版本**"):
            i += 1
            continue
        if re.match(r"^(?:- |\d+\. )", line):
            number = re.match(r"^(\d+)\. ", line)
            content = re.sub(r"^(?:- |\d+\. )", "", line)
            p = doc.add_paragraph(style="Normal" if number else "List Bullet")
            if number:
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                content = f"{number.group(1)}. {content}"
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, content)
            i += 1
            continue
        parts = [line.rstrip("  ")]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(?:#|\||```|---|- |\d+\. )", lines[i].strip()):
            parts.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(parts))
    count = append_source_index(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"source_lines={len(lines)} tables={tables} diagrams={diagrams} source_index_items={count}")


if __name__ == "__main__":
    build()
