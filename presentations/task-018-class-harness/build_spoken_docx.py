"""Package the approved oral text as a readable Word manuscript."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT/"spoken-script.md"
OUTPUT=ROOT/"Harness_课堂汇报_完整讲稿.docx"
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0)
sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
for name,size in [("Normal",11),("Title",19),("Heading 1",13.5)]:
    style=doc.styles[name]
    style.font.name="Microsoft YaHei"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string("000000")
    style.font.italic=False
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=1.25
doc.styles["Title"].font.bold=True
doc.styles["Title"].paragraph_format.space_after=Pt(8)
doc.styles["Heading 1"].font.bold=True
doc.styles["Heading 1"].paragraph_format.keep_with_next=True
doc.styles["Heading 1"].paragraph_format.space_before=Pt(12)
doc.styles["Heading 1"].paragraph_format.space_after=Pt(4)
for style in doc.styles:
    for border in list(style._element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
head=sec.header.paragraphs[0]
head.add_run("Agent Harness 课堂汇报讲稿").font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run("第 ")
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
foot._p.append(field);foot.add_run(" 页")
blocks=SOURCE.read_text(encoding="utf-8").strip().split("\n\n")
slide=0
for text in blocks:
    if text.startswith("# "):
        doc.add_paragraph(text[2:],style="Title")
    elif text.startswith("## "):
        slide+=1
        p=doc.add_paragraph(text[3:],style="Heading 1")
        if slide in [3,5,7]:p.paragraph_format.page_break_before=True
    elif text.startswith("建议用时"):
        p=doc.add_paragraph(text);p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(7)
        for run in p.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string("666666")
    else:
        p=doc.add_paragraph(text)
        p.paragraph_format.widow_control=True
        if slide==0:
            for run in p.runs:run.font.size=Pt(9.5);run.font.color.rgb=RGBColor.from_string("555555")
doc.core_properties.title="光子芯片 Agent Harness 课堂汇报讲稿"
doc.core_properties.subject="八页项目重点版口头汇报"
doc.core_properties.author="PIC_Harness_Research"
assert slide==8
doc.save(OUTPUT)
back=Document(OUTPUT)
expected=[re.sub(r"^#{1,2} ","",b) for b in blocks]
actual=[p.text for p in back.paragraphs if p.text]
assert expected==actual
print(f"Wrote {OUTPUT}; 8 sections; {len(actual)} paragraphs; exact text readback passed.")
