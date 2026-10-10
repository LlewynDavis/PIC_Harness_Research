"""Adapt two native slides from the retained template, using the requested source deck."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parent
A=ROOT/"assets";A.mkdir(exist_ok=True)
parser=argparse.ArgumentParser()
parser.add_argument("--source",required=True,type=Path)
parser.add_argument("--reference",type=Path,default=Path.home()/".codex/skills/artifact-template-engineering-optics/assets/reference.pptx")
args=parser.parse_args()
plan=json.loads((ROOT/"outline.json").read_text(encoding="utf-8"))
source=Presentation(args.source)
assert len(source.slides)>=7
assert "Harness" in source.slides[0].shapes[0].text
src_pic=next(sh for sh in source.slides[6].shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE and sh.name=="PaperOriginal_Kharel_Fig2a")
image=A/f"kharel-fig2a{src_pic.image.ext if src_pic.image.ext.startswith('.') else '.'+src_pic.image.ext}"
image.write_bytes(src_pic.image.blob)
source_text="\n".join(s.text for s in source.slides[6].shapes if s.has_text_frame)
for value in ["0.460","0.109","400","1550","150"]:assert value in source_text,value

prs=Presentation(args.reference)
template_sha=hashlib.sha256(args.reference.read_bytes()).hexdigest()
assert template_sha=="6cdcf2d54733b902d606a0c9719a3ac212237cf00fb29859a54541328c9a4332"
keep={3,5}
for idx in reversed(range(len(prs.slides))):
    if idx not in keep:
        entry=prs.slides._sldIdLst[idx]
        prs.part.drop_rel(entry.rId)
        prs.slides._sldIdLst.remove(entry)
TEAL="006579";GREY="5A5A5A";INK="383838"

def format_text(sh,text,size=16,color=INK,bold=False,align=PP_ALIGN.LEFT):
    tf=sh.text_frame;tf.clear()
    tf.margin_left=Inches(.08);tf.margin_right=Inches(.08)
    tf.margin_top=Inches(.03);tf.margin_bottom=Inches(.03)
    tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4)
        p.line_spacing=1.15
        r=p.add_run();r.text=line;r.font.name="微软雅黑"
        r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        r._r.get_or_add_rPr().set("lang","zh-CN")
        ea=OxmlElement("a:ea");ea.set("typeface","微软雅黑");r._r.get_or_add_rPr().append(ea)
    return sh
def at(sh,x,y,w,h):
    sh.left=Inches(x);sh.top=Inches(y);sh.width=Inches(w);sh.height=Inches(h)
    return sh
def clone(slide,element,name):
    el=copy.deepcopy(element)
    c=el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    c.set("id",str(max([s.shape_id for s in slide.shapes]+[1])+1));c.set("name",name)
    slide.shapes._spTree.insert_element_before(el,"p:extLst")
    return next(s for s in slide.shapes if s._element is el)
def textbox(slide,text,x,y,w,h,size=16,color=INK,bold=False,align=PP_ALIGN.LEFT,name="EditableText"):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=name
    return format_text(sh,text,size,color,bold,align)
def connector(slide,x1,y1,x2,y2,arrow=False):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=RGBColor.from_string(TEAL);sh.line.width=Pt(1.6)
    if arrow:
        ln=sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
        end=OxmlElement("a:tailEnd");end.set("type","triangle");end.set("w","med");end.set("len","med");ln.append(end)
    return sh

# Change native navigation labels only on the copied layouts.
updated=set()
for s in prs.slides:
    layout=s.slide_layout
    if layout.part.partname in updated:continue
    updated.add(layout.part.partname)
    nav=sorted([sh for sh in layout.shapes if sh.has_text_frame and sh.top<Inches(.8) and sh.width>Inches(2)],key=lambda sh:sh.left)
    assert len(nav)==4
    for sh,label in zip(nav,plan["navigation"]):format_text(sh,label,22,GREY,False,PP_ALIGN.CENTER)

for i,(s,spec) in enumerate(zip(prs.slides,plan["slides"])):
    s._element.set("showMasterSp","0")
    for tag in ["{http://schemas.openxmlformats.org/presentationml/2006/main}timing", "{http://schemas.openxmlformats.org/presentationml/2006/main}transition"]:
        for stale in list(s._element.iterchildren(tag)):
            s._element.remove(stale)
    for alternate in list(s._element):
        if alternate.tag.endswith("AlternateContent") and any(el.tag.endswith("}transition") for el in alternate.iter()):
            s._element.remove(alternate)
    active=next(sh for sh in s.shapes if sh.has_text_frame and sh.top<Inches(.8) and sh.width>Inches(2))
    nav=sorted([sh for sh in s.slide_layout.shapes if sh.has_text_frame and sh.top<Inches(.8) and sh.width>Inches(2)],key=lambda sh:sh.left)
    target=nav[spec["active_navigation"]]
    active_el=copy.deepcopy(active._element)
    active.left=target.left;active.top=target.top;active.width=target.width;active.height=target.height
    format_text(active,plan["navigation"][spec["active_navigation"]],22,"FFFFFF",True,PP_ALIGN.CENTER)
    if i==0:
        backgrounds=[copy.deepcopy(sh._element) for sh in s.shapes if sh.name in ["矩形: 圆角 17","矩形: 圆角 18","矩形: 圆角 19"]]
        label_el=copy.deepcopy(next(sh for sh in s.shapes if sh.has_text_frame and sh.text=="点式感知")._element)
        title_el=copy.deepcopy(next(sh for sh in s.shapes if sh.has_text_frame and "光纤如何" in sh.text)._element)
        assert len(backgrounds)==3
    else:
        frame_el=copy.deepcopy(s.shapes[0]._element)
        desc_el=copy.deepcopy(next(sh for sh in s.shapes if sh.has_text_frame and "电力电缆" in sh.text)._element)
    for sh in list(s.shapes):
        if sh is not active and sh._element is not active._element:
            sh._element.getparent().remove(sh._element)
    active._element.getparent().remove(active._element)
    for j,(nav_source,label) in enumerate(zip(nav,plan["navigation"])):
        head=clone(s,active_el if j==spec["active_navigation"] else nav_source._element,f"Editable_Navigation_{j+1}")
        head.left=nav_source.left;head.top=nav_source.top;head.width=nav_source.width;head.height=nav_source.height
        head.fill.solid();head.fill.fore_color.rgb=RGBColor.from_string(TEAL if j==spec["active_navigation"] else "FFFFFF")
        format_text(head,label,22,"FFFFFF" if j==spec["active_navigation"] else GREY,j==spec["active_navigation"],PP_ALIGN.CENTER)
    logo_source=next(sh for sh in s.slide_layout.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE)
    logo_path=A/f"template-logo.{logo_source.image.ext}"
    logo_path.write_bytes(logo_source.image.blob)
    logo=s.shapes.add_picture(str(logo_path),logo_source.left,logo_source.top,logo_source.width,logo_source.height)
    logo.name="Template_Logo"
    logo.crop_left=logo_source.crop_left;logo.crop_right=logo_source.crop_right
    logo.crop_top=logo_source.crop_top;logo.crop_bottom=logo_source.crop_bottom
    logo.rotation=logo_source.rotation
    if i==0:
        title=clone(s,title_el,"Editable_Principle_Title")
        at(title,.55,1.10,12.2,.68);format_text(title,spec["title"],26,INK,True)
        textbox(s,spec["subtitle"],.65,1.92,12.0,.44,16,GREY,name="Editable_Principle_Subtitle")
        for j,(label,body) in enumerate(spec["modules"]):
            x=.65+j*3.16
            bg=clone(s,backgrounds[min(j,2)],f"Editable_Module_{j+1}_Frame")
            at(bg,x,2.67,2.55,2.05)
            badge=clone(s,label_el,f"Editable_Module_{j+1}_Label")
            at(badge,x+.14,2.82,2.27,.48);format_text(badge,label,18,"FFFFFF",True,PP_ALIGN.CENTER)
            textbox(s,"\n".join(body),x+.12,3.65,2.31,.75,16,GREY,align=PP_ALIGN.CENTER,name=f"Editable_Module_{j+1}_Body")
            if j<3:connector(s,x+2.61,3.69,x+3.08,3.69,True)
        connector(s,11.405,4.81,11.405,5.28)
        connector(s,11.405,5.28,5.085,5.28)
        connector(s,5.085,5.28,5.085,4.81,True)
        textbox(s,spec["feedback"],6.0,5.40,6.15,.40,18,TEAL,True,PP_ALIGN.CENTER,"Editable_Feedback")
        textbox(s,spec["trace"],.75,5.40,3.8,.40,16,GREY,name="Editable_Trace")
        strip=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.2),Inches(6.25),Inches(11.45),Inches(.64))
        strip.name="Editable_Takeaway_Band";strip.fill.solid();strip.fill.fore_color.rgb=RGBColor.from_string(TEAL);strip.line.fill.background()
        format_text(strip,spec["takeaway"],18,"FFFFFF",True,PP_ALIGN.CENTER)
    else:
        textbox(s,spec["title"],.65,1.00,12.0,.58,26,TEAL,True,name="Editable_Example_Title")
        textbox(s,spec["subtitle"],.65,1.64,12.0,.40,16,GREY,name="Editable_Example_Subtitle")
        frame=clone(s,frame_el,"Paper_Plot_Frame")
        at(frame,.65,2.18,7.35,4.59)
        textbox(s,spec["plot_title"],.9,2.29,6.85,.40,18,TEAL,True,PP_ALIGN.CENTER,"Editable_Plot_Title")
        iw,ih=src_pic.image.size
        width=6.96;height=width*ih/iw
        pic=s.shapes.add_picture(str(image),Inches(.84),Inches(2.91),width=Inches(width),height=Inches(height))
        pic.name="PaperOriginal_Kharel_Fig2a"
        connector(s,8.06,4.17,8.48,4.17,True)
        desc=clone(s,desc_el,"Editable_Example_Description")
        at(desc,8.62,2.29,4.00,.45);format_text(desc,spec["metric_label"],22,INK,True)
        textbox(s,spec["metric"],8.62,2.89,4.00,.60,27,TEAL,True,name="Editable_Result")
        textbox(s,spec["metric_note"],8.62,3.49,4.00,.38,15,GREY,name="Editable_Metric_Boundary")
        textbox(s,spec["conditions_title"],8.62,4.05,4.00,.40,18,TEAL,True,name="Editable_Conditions_Title")
        textbox(s,"\n".join(spec["conditions"]),8.62,4.54,4.00,1.09,16,GREY,name="Editable_Conditions")
        textbox(s,"\n".join(spec["interpretation"]),8.62,5.78,4.00,1.12,14,GREY,name="Editable_Interpretation")
    textbox(s,spec["source"],1.1,7.06,11.6,.24,9,GREY,name="Editable_Source")
    s.notes_slide.notes_text_frame.text=spec["notes"]

output=ROOT/"Harness_课题介绍_原理与案例_模板版.pptx"
prs.core_properties.title=plan["title"]
prs.core_properties.subject="两页课题组介绍模块 原理示意与文献仿真案例"
prs.core_properties.author="PIC_Harness_Research"
prs.save(output)
manifest={"source_filename":args.source.name,"source_sha256":hashlib.sha256(args.source.read_bytes()).hexdigest(),"source_pages":[1,7],"template_sha256":template_sha,"template_pages":[4,6],"image_sha256":hashlib.sha256(src_pic.image.blob).hexdigest(),"canvas_inches":[prs.slide_width/914400,prs.slide_height/914400],"nav_note":"四个导航标题为模块的可编辑通用标签，实际插入课题组PPT后可按其章节改名。","output":output.name}
(ROOT/"source-manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("Adapted 2 native slides from template pages 4/6; original case figure preserved.")
