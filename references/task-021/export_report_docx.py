"""Create an editable Word copy of the current complete report, without rewriting it."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path
from urllib.parse import unquote

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).parent / "chapters/完整调研报告.md"
OUTPUT = ROOT / "output/task-021/TASK-021_完整调研报告_2026-10-16.docx"
WIDTH_PT = 507.28


def set_font(element, size=11, bold=False, family="宋体", color="000000"):
    if hasattr(element, "font"):
        element.font.name = family
        element.font.size = Pt(size)
        element.font.bold = bold
        element.font.color.rgb = RGBColor.from_string(color)
    rpr = element._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), family)


def inline(paragraph, text, size=11, bold=False):
    for part in re.split(r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)", text):
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if match:
            label, target = match.groups()
            if not re.match(r"https?://", target):
                filename, _, fragment = target.partition("#")
                target = (SOURCE.parent / unquote(filename)).resolve().as_uri()
                if fragment:
                    target += "#" + fragment
            rel = paragraph.part.relate_to(
                target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            for key in ("ascii", "hAnsi", "eastAsia"):
                fonts.set(qn("w:" + key), "宋体")
            rpr.append(fonts)
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(round(size * 2))); rpr.append(sz)
            color = OxmlElement("w:color"); color.set(qn("w:val"), "176C73"); rpr.append(color)
            if bold:
                rpr.append(OxmlElement("w:b"))
            run.append(rpr)
            node = OxmlElement("w:t"); node.text = label; run.append(node)
            link.append(run); paragraph._p.append(link)
        else:
            is_bold = part.startswith("**") and part.endswith("**")
            if is_bold:
                part = part[2:-2]
            elif part.startswith("`") and part.endswith("`"):
                part = part[1:-1]
            set_font(paragraph.add_run(part), size, bold or is_bold)


def make_table(doc, lines):
    rows = []
    for line in lines:
        cells = [s.strip() for s in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    n = len(rows[0])
    ratios = {3: [.22, .39, .39], 4: [.16, .28, .28, .28]}.get(n, [1 / n] * n)
    table = doc.add_table(rows=0, cols=n)
    table.autofit = False
    for col, ratio in zip(table.columns, ratios):
        col.width = Pt(WIDTH_PT * ratio)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, value, ratio in zip(row.cells, values, ratios):
            cell.width = Pt(WIDTH_PT * ratio)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "start", "bottom", "end"):
                m = OxmlElement("w:" + side)
                m.set(qn("w:w"), "85"); m.set(qn("w:type"), "dxa"); margins.append(m)
            pr.append(margins)
            if row_index == 0:
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "E5EFF1"); pr.append(shd)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = Pt(14)
            p.paragraph_format.keep_with_next = row_index == 0
            inline(p, value, 9.5, row_index == 0)


def main():
    original = SOURCE.read_bytes()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.left_margin = Pt(44); sec.right_margin = Pt(44)
    sec.top_margin = Pt(43); sec.bottom_margin = Pt(45)
    sec.footer_distance = Pt(20)
    normal = doc.styles["Normal"]
    set_font(normal)
    pf = normal.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.first_line_indent = Pt(22); pf.line_spacing = Pt(18)
    pf.widow_control = True
    for name, size in (("Title", 19), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        set_font(style, size, True, "黑体")
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = Pt(size + 8)
    # Remove inherited template rules: the report title has no decorative line.
    for style in doc.styles:
        for node in list(style.element.iter(qn("w:pBdr"))):
            node.getparent().remove(node)
    doc.core_properties.title = "完整调研报告"
    doc.core_properties.author = "PIC_Harness_Research"
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

    lines = original.decode("utf-8-sig").splitlines()
    i = 0
    tables = figures = 0
    references = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            make_table(doc, block); tables += 1; continue
        heading = re.match(r"^(#{1,3})\s+(.+)", line)
        if heading:
            level, title = len(heading[1]), heading[2]
            p = doc.add_paragraph(style="Title" if not doc.paragraphs else f"Heading {level}")
            # Preserve every heading verbatim, including the author's punctuation.
            p.add_run(title)
            references = title == "参考文献与来源" or references
            i += 1; continue
        picture = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if picture:
            asset = (SOURCE.parent / picture[2]).resolve()
            if asset.suffix.lower() == ".svg":
                asset = asset.with_suffix(".png")
            with Image.open(asset) as image:
                width, height = image.size
            scale = min(WIDTH_PT / width, (320 if "fig2-original" in asset.name else 290) / height)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(5)
            # Exact body leading would clip inline images to a single text line.
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = True
            pic = p.add_run().add_picture(str(asset), width=Pt(width * scale), height=Pt(height * scale))
            pic._inline.docPr.set("descr", picture[1])
            figures += 1; i += 1; continue
        caption = bool(re.match(r"图\s*\d", line))
        p = doc.add_paragraph()
        if caption:
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = Pt(13)
            p.paragraph_format.space_after = Pt(4)
        if re.match(r"^2026-\d\d-\d\d", line):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
        if line.startswith("> "):
            line = line[2:]
        line = re.sub(r"([1-9])[\ufe0f\u20e3]+", r"\1. ", line)
        inline(p, line, 9 if caption else 10 if references else 11)
        i += 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    assert SOURCE.read_bytes() == original
    print("Output:", OUTPUT)
    print("Tables:", tables, "Figures:", figures)
    print("Source SHA256:", hashlib.sha256(original).hexdigest())


if __name__ == "__main__":
    main()
